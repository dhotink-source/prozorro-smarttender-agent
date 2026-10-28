import os
import io
import re
import sys
import time
import html
import json
import zipfile
import sqlite3
import subprocess
import tempfile
from datetime import datetime, timedelta, timezone
import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
import docx
from concurrent.futures import ThreadPoolExecutor, as_completed
from google import genai
from google.genai import types

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(line_buffering=True)

# 1. Конфигурация ключей из GitHub Secrets
GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "").strip()
GOOGLE_WEB_APP_URL = os.environ.get("GOOGLE_WEB_APP_URL")
SECRET_BRIDGE_KEY = os.environ.get("SECRET_BRIDGE_KEY", "FUEL_TENDER_SECRET_2026")

PROZORRO_API_BASE = "https://public-api.prozorro.gov.ua/api/2.5"
MIN_AMOUNT_UAH = 2400000

# Поиск в Prozorro идет строго по целевым кодам ДК 021:2015
TARGET_CPV_PREFIXES = ("09130000", "09132", "09134")
EXCLUDED_CPV_PREFIXES = ("09131", "09133", "09135")
MAX_GEMINI_ANALYSES_PER_RUN = 15

GEMINI_MODELS_FALLBACK = (
    "gemini-3-flash-preview",
    "gemini-3.0-flash",
    "gemini-3.1-pro-preview",
    "gemini-3.1-flash-lite-preview",
    "gemini-2.5-flash",
)

EXCLUDED_FUEL_REGEX = re.compile(
    r"(\bгас\b|\bгасу\b|гасов[еоиій]|газ\b|\bгазу\b|газов[еоиій]|скраплен|пропан|бутан|метан|"
    r"авіаційн|реактивн|газотурбін|jet\s*a|джет|\bрт\b|\bтс-?1\b|мазут|олив[аи]|мастил|пічн[ео]\s+палив)",
    re.IGNORECASE
)

MANDATORY_FUEL_REGEX = re.compile(
    r"(дизельн|паливо\s+дизельн|пальне\s+дизельн|\bдп\b|дп-|бензин|\bа-?95\b|\bа-?92\b|\bа-?98\b|\bа-?80\b)",
    re.IGNORECASE
)

COUPON_CARD_REGEX = re.compile(
    r"(талон|скретч|смарт-?карт|паливн[іи]х?\s+карт|бланк[іи]?в?-дозвол|заправк[аи]\s+на\s+азс|відпуск\s+з\s+азс)",
    re.IGNORECASE
)
BULK_KEYWORD_REGEX = re.compile(r"(налив|автоцистерн|бензовоз|вагон-?цистерн)", re.IGNORECASE)

ZERO_VAT_REGEX = re.compile(
    r"(постанов[аиуою]\s+кму\s*№?\s*178|від\s+02\.03\.2022\s*№?\s*178|пдв\s*0\s*%|0\s*%\s*пдв|"
    r"нульов[оа]ю\s+ставк[оа]ю|за\s+ставкою\s+0\s*%)",
    re.IGNORECASE
)

# Регулярные выражения для подстраховки поиска % превышения бюджета или запрета в ТД
OVER_BUDGET_PERCENT_REGEX = re.compile(
    r"відсоток\s+перевищення\s+ціни\s+тендерної\s+пропозиції[^.]{0,250}?не\s+більше\s+(?:ніж\s+)?(\d+(?:[.,]\d+)?)\s*%",
    re.IGNORECASE | re.DOTALL
)
OVER_BUDGET_PROHIBITED_REGEX = re.compile(
    r"(не\s+приймає(?:ться)?\s+до\s+розгляду\s+тендерн[ау]\s+пропозиці[яю],\s+ціна\s+якої\s+є\s+вищою|"
    r"ціна\s+тендерної\s+пропозиції\s+не\s+може\s+перевищувати\s+очікувану\s+вартість|"
    r"відхиляє\s+тендерну\s+пропозицію[^.]{0,120}?ціна\s+якої\s+перевищує\s+очікувану\s+вартість)",
    re.IGNORECASE | re.DOTALL
)

# Очистка отображаемого текста пункта 1 от кодов ДК 021:2015 (при сохранении самого фильтра поиска по ДК)
CPV_STRIP_REGEX = re.compile(
    r"(?:,?\s*\(?(?:код\s*)?(?:за\s*)?ДК\s*021:2015\s*[:\-–—]?\s*\d{8}-\d[^,.)]*\)?|,?\s*\b0913\d{4}-\d\b)",
    re.IGNORECASE
)

PROCUREMENT_METHODS = {
    "aboveThreshold": "Открытые торги с особенностями",
    "aboveThresholdUA": "Открытые торги",
    "aboveThresholdEU": "Открытые торги с публикацией на английском языке",
    "priceQuotation": "Запрос ценовых предложений",
    "negotiation": "Переговорная процедура",
    "negotiation.quick": "Переговорная процедура (ускоренная)",
    "belowThreshold": "Упрощенная закупка",
    "reporting": "Отчет о заключенном договоре"
}

http_session = requests.Session()
retries = Retry(total=3, backoff_factor=0.5, status_forcelist=[500, 502, 503, 504])
adapter = HTTPAdapter(pool_connections=20, pool_maxsize=20, max_retries=retries)
http_session.mount("https://", adapter)
http_session.mount("http://", adapter)

conn = sqlite3.connect("processed_tenders.db")
cursor = conn.cursor()
cursor.execute("CREATE TABLE IF NOT EXISTS tenders (id TEXT PRIMARY KEY)")
cursor.execute("CREATE TABLE IF NOT EXISTS sync_state (key TEXT PRIMARY KEY, val TEXT)")
cursor.execute("CREATE TABLE IF NOT EXISTS retry_queue (id TEXT PRIMARY KEY)")
conn.commit()

vertex_clients = []
if GEMINI_API_KEY.startswith("{"):
    creds_info = json.loads(GEMINI_API_KEY)
    project_id = creds_info.get("project_id", "ultra-dominion-510720-s6")
    if creds_info.get("type") == "service_account":
        from google.oauth2 import service_account
        creds = service_account.Credentials.from_service_account_info(
            creds_info,
            scopes=["https://www.googleapis.com/auth/cloud-platform"]
        )
    else:
        from google.oauth2.credentials import Credentials
        creds = Credentials.from_authorized_user_info(
            creds_info,
            scopes=["https://www.googleapis.com/auth/cloud-platform"]
        )
    for loc in ("global", "us-central1"):
        vertex_clients.append((
            loc,
            genai.Client(vertexai=True, project=project_id, location=loc, credentials=creds)
        ))
else:
    vertex_clients.append(("ai_studio", genai.Client(api_key=GEMINI_API_KEY)))

def is_already_processed(tender_id: str) -> bool:
    cursor.execute("SELECT 1 FROM tenders WHERE id = ?", (tender_id,))
    return cursor.fetchone() is not None

def mark_as_processed(tender_id: str):
    cursor.execute("INSERT OR IGNORE INTO tenders (id) VALUES (?)", (tender_id,))
    cursor.execute("DELETE FROM retry_queue WHERE id = ?", (tender_id,))
    conn.commit()

def add_to_retry_queue(tender_id: str):
    cursor.execute("INSERT OR IGNORE INTO retry_queue (id) VALUES (?)", (tender_id,))
    conn.commit()

def get_retry_queue_ids() -> list:
    cursor.execute("SELECT id FROM retry_queue")
    return [r[0] for r in cursor.fetchall()]

def get_sync_val(key: str) -> str:
    cursor.execute("SELECT val FROM sync_state WHERE key = ?", (key,))
    row = cursor.fetchone()
    return row[0] if row else ""

def set_sync_val(key: str, val: str):
    cursor.execute("INSERT OR REPLACE INTO sync_state (key, val) VALUES (?, ?)", (key, val))
    conn.commit()

def get_total_processed_count() -> int:
    cursor.execute("SELECT COUNT(*) FROM tenders")
    row = cursor.fetchone()
    return int(row[0]) if row else 0

# Проверка номенклатуры по целевым кодам ДК и ключевым словам топлива
def is_valid_diesel_or_gasoline_tender(tender: dict) -> tuple[bool, str]:
    title_and_desc = f"{tender.get('title', '')} {tender.get('description', '')}"
    items = tender.get("items", [])

    valid_fuel_items = []
    for item in items:
        cpv = str(item.get("classification", {}).get("id", ""))
        item_desc = str(item.get("description", ""))
        if cpv.startswith(EXCLUDED_CPV_PREFIXES):
            continue
        if not cpv.startswith(TARGET_CPV_PREFIXES):
            continue
        if EXCLUDED_FUEL_REGEX.search(item_desc) and not MANDATORY_FUEL_REGEX.search(item_desc):
            continue
        valid_fuel_items.append(item_desc)

    if not valid_fuel_items:
        return False, "Нет позиций бензина или ДТ по целевым кодам ДК"

    combined_items_text = " ".join(valid_fuel_items)
    full_check_text = f"{title_and_desc} {combined_items_text}"

    if EXCLUDED_FUEL_REGEX.search(full_check_text) and not MANDATORY_FUEL_REGEX.search(combined_items_text):
        return False, "Отсеян по стоп-словам"

    if not MANDATORY_FUEL_REGEX.search(full_check_text):
        return False, "Не найдено упоминаний бензина или дизельного топлива"

    return True, "OK"

def is_strictly_coupons_or_cards(tender: dict) -> bool:
    full_header = f"{tender.get('title', '')} {tender.get('description', '')}"
    items_desc = " ".join(str(i.get("description", "")) for i in tender.get("items", []))
    combined_text = f"{full_header} {items_desc}"

    if BULK_KEYWORD_REGEX.search(combined_text):
        return False

    lots = tender.get("lots", [])
    if len(lots) > 1:
        active_lots = [l for l in lots if l.get("status", "active") == "active"]
        if active_lots:
            return all(
                bool(COUPON_CARD_REGEX.search(f"{l.get('title', '')} {l.get('description', '')}"))
                for l in active_lots
            )

    return bool(COUPON_CARD_REGEX.search(full_header))

def build_smarttender_url(tender_ua_id: str) -> str:
    gateway_url = f"https://smarttender.biz/government/details/{tender_ua_id}"
    try:
        resp = http_session.head(gateway_url, allow_redirects=True, timeout=8)
        if resp.status_code == 200 and "smarttender.biz" in resp.url:
            return resp.url
    except Exception:
        pass
    return gateway_url

def format_deadline_strict(iso_str: str) -> str:
    if not iso_str:
        return "Не указан"
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})T(\d{2}:\d{2})", str(iso_str))
    if m:
        year, month, day, time_hm = m.group(1), m.group(2), m.group(3), m.group(4)
        return f"{time_hm}, {day}-{month}-{year}"
    return str(iso_str)[:16].replace("T", " ")

def format_date_only(iso_str: str) -> str:
    if not iso_str:
        return ""
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", str(iso_str))
    if m:
        return f"{m.group(3)}.{m.group(2)}.{m.group(1)}"
    return str(iso_str)[:10]

def parse_iso_dt(iso_str: str):
    if not iso_str:
        return None
    try:
        clean = str(iso_str).replace("Z", "+00:00")
        return datetime.fromisoformat(clean)
    except Exception:
        return None

# Очистка отображаемого текста от кодов ДК 021:2015, англицизмов (Milestones/Items) и добавление источника к пунктам
def clean_and_enforce_source(text: str, default_source: str = "Договора", strip_cpv: bool = False) -> str:
    if not text:
        return ""
    cleaned = str(text)
    if strip_cpv:
        cleaned = CPV_STRIP_REGEX.sub("", cleaned)
        cleaned = re.sub(r"\s{2,}", " ", cleaned).strip(" ,.")

    cleaned = re.sub(r"\bmilestones\b", "Электронных условий объявления", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\b(items|номенклатура\s+prozorro)\b", "Номенклатуры объявления", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\blots\b", "Лотов объявления", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\btenderPeriod\b", "Сроков подачи объявления", cleaned, flags=re.IGNORECASE)

    pattern = re.compile(r"\((п\.?\s*[\d\.,\s\-а-яіїєґА-ЯІЇЄҐ]+)\)")
    source_keywords = ("догов", "контракт", "тд", "документац", "додат", "прилож", "специф", "оголош", "закон", "постанов", "услов", "номенкл")

    def replacer(match):
        inner = match.group(1).strip()
        low_inner = inner.lower()
        if any(k in low_inner for k in source_keywords):
            return f"({inner})"
        return f"({inner} {default_source})"

    return pattern.sub(replacer, cleaned)

def format_percent_clean(val: float) -> str:
    if abs(val - round(val)) < 1e-6:
        return str(int(round(val)))
    return str(round(val, 2)).rstrip("0").rstrip(".")

# Расчет бюджета с НДС и проверкой возможности подачи предложений сверх ожидаемой стоимости (4 сценария)
def extract_pricing_meta_and_budget(tender_data: dict, ai: dict, docs_text: str) -> dict:
    val = tender_data.get("value", {})
    budget_raw = float(val.get("amount", 0) or 0)
    prozorro_vat_included = bool(val.get("valueAddedTaxIncluded", False))
    method_type = tender_data.get("procurementMethodType", "")

    ai_vat_rate = ai.get("vat_rate_percent", 20)
    is_zero_vat = (ai_vat_rate == 0) or bool(ZERO_VAT_REGEX.search(docs_text))
    vat_rate_num = 0 if is_zero_vat else 20

    if is_zero_vat:
        multiplier = 1.0
        vat_label = "с НДС (ставка НДС: 0%)"
    else:
        multiplier = 1.0 if prozorro_vat_included else 1.2
        vat_label = "с НДС (ставка НДС: 20%)"

    items = tender_data.get("items", [])
    fuel_items = [
        i for i in items
        if str(i.get("classification", {}).get("id", "")).startswith(TARGET_CPV_PREFIXES)
        and not str(i.get("classification", {}).get("id", "")).startswith(EXCLUDED_CPV_PREFIXES)
    ]
    target_items = fuel_items if fuel_items else items
    items_count = len(target_items)
    single_qty = 0.0
    unit_label = "ед."

    is_single_item = False
    base_display_num = 0.0

    if items_count == 1:
        single_qty = float(target_items[0].get("quantity", 0) or 0)
        unit_raw = str(target_items[0].get("unit", {}).get("name", "ед.")).lower()
        if "літр" in unit_raw or "литр" in unit_raw or unit_raw == "л":
            unit_label = "литр"
        elif "тон" in unit_raw or unit_raw == "т":
            unit_label = "тонну"
        elif "кілогр" in unit_raw or "килогр" in unit_raw or unit_raw == "кг":
            unit_label = "кг"
        else:
            unit_label = unit_raw

        if single_qty > 0:
            is_single_item = True
            base_display_num = (budget_raw / single_qty) * multiplier
            base_budget_str = f"{base_display_num:,.2f} грн/{unit_label} {vat_label}".replace(",", " ")
        else:
            base_display_num = budget_raw * multiplier
            base_budget_str = f"{base_display_num:,.2f} грн {vat_label}".replace(",", " ")
    else:
        base_display_num = budget_raw * multiplier
        base_budget_str = f"{base_display_num:,.2f} грн {vat_label}".replace(",", " ")

    # Определение политики превышения бюджета («Ціна пропозиції обмежена очікуваною вартістю: Ні / Так»)
    ai_policy = str(ai.get("over_budget_policy", "") or "").strip().upper()
    try:
        over_pct = float(ai.get("over_budget_percent", 0) or 0)
    except Exception:
        over_pct = 0.0

    # Дополнительная проверка регулярным выражением по тексту ТД (на случай, если ИИ пропустил цифру)
    if over_pct <= 0 and docs_text:
        m_pct = OVER_BUDGET_PERCENT_REGEX.search(docs_text)
        if m_pct:
            try:
                over_pct = float(m_pct.group(1).replace(",", "."))
                ai_policy = "ALLOWED_WITH_PERCENT"
            except Exception:
                pass
        elif OVER_BUDGET_PROHIBITED_REGEX.search(docs_text) and method_type == "aboveThreshold":
            ai_policy = "PROHIBITED_IN_TD"

    # В процедуре «Открытые торги с особенностями» (aboveThreshold) параметр в объявлении Prozorro/SmartTender всегда «Ні»
    announcement_allows_over_budget = (method_type == "aboveThreshold") or (over_pct > 0)

    over_budget_suffix = ""
    if announcement_allows_over_budget:
        if over_pct > 0 or ai_policy == "ALLOWED_WITH_PERCENT":
            pct_clean = format_percent_clean(over_pct)
            real_val = base_display_num * (1.0 + over_pct / 100.0)
            if is_single_item:
                real_str = f"{real_val:,.2f} грн/{unit_label} с НДС".replace(",", " ")
            else:
                real_str = f"{real_val:,.2f} грн с НДС".replace(",", " ")
            over_budget_suffix = (
                f". Есть возможность подачи предложений сверх бюджета на +{pct_clean}%, "
                f"то есть реальный бюджет закупки {real_str}"
            )
        elif ai_policy == "PROHIBITED_IN_TD":
            over_budget_suffix = (
                ". Внимание! Объявлением предусмотрена возможность подачи предложений сверх бюджета, "
                "однако тендерной документацией эта возможность запрещена. Ошибка заказчика!"
            )
        else:
            over_budget_suffix = (
                ". Предусмотрена возможность подачи предложений сверх бюджета, "
                "однако величина превышения бюджета в закупке не установлена."
            )

    full_budget_str = base_budget_str + over_budget_suffix
    return {
        "budgetDisplay": full_budget_str,
        "isZeroVat": is_zero_vat,
        "vatRatePercent": vat_rate_num,
        "itemsCount": 1 if is_single_item else items_count,
        "singleItemQty": single_qty if is_single_item else 0.0,
        "unitLabel": unit_label
    }

def format_participant_price_with_vat(
    raw_amount: float,
    bid_vat_included: bool,
    is_zero_vat: bool,
    items_count: int,
    single_qty: float,
    unit_label: str
) -> str:
    if is_zero_vat:
        multiplier = 1.0
        vat_label = "с НДС (ставка НДС: 0%)"
    else:
        multiplier = 1.0 if bid_vat_included else 1.2
        vat_label = "с НДС (ставка НДС: 20%)"

    total_with_vat = float(raw_amount or 0) * multiplier
    if items_count == 1 and single_qty > 0:
        unit_price = total_with_vat / single_qty
        return f"{unit_price:,.2f} грн/{unit_label} {vat_label}".replace(",", " ")
    return f"{total_with_vat:,.2f} грн {vat_label}".replace(",", " ")

def next_business_day(dt: datetime, days_to_add: int = 1) -> datetime:
    cur = dt
    added = 0
    while added < days_to_add:
        cur += timedelta(days=1)
        if cur.weekday() < 5:
            added += 1
    return cur

# Предварительный математический расчет дат по законодательству для передачи в Gemini
def build_statutory_dates_reference(tender_data: dict) -> str:
    method_type = tender_data.get("procurementMethodType", "")
    end_iso = tender_data.get("tenderPeriod", {}).get("endDate", "")
    dt_end = parse_iso_dt(end_iso) or datetime.now(timezone(timedelta(hours=2)))
    d_str = dt_end.strftime("%d.%m.%Y (%H:%M)")
    next_bd = next_business_day(dt_end, 1)
    next_bd_str = next_bd.strftime("%d.%m.%Y")

    if method_type == "priceQuotation":
        return (
            f"Процедура: Запрос ценовых предложений (Постановление КМУ №822 и №1178). "
            f"Окончание подачи предложений (День D): {d_str}. "
            f"Аукцион и мораторий на обжалование (5 дней) законом НЕ предусмотрены. "
            f"Раскрытие предложений происходит автоматически в День D ({dt_end.strftime('%d.%m.%Y')}). "
            f"При оперативности заказчика решение об определении победителя и подписание договора КЭП возможны "
            f"уже в день раскрытия предложений ({dt_end.strftime('%d.%m.%Y')}) или на следующий рабочий день ({next_bd_str})."
        )
    elif method_type == "belowThreshold":
        return (
            f"Процедура: Упрощенная закупка (ст. 14 Закона Украины 'О публичных закупках'). "
            f"Окончание подачи (День D): {d_str}. Мораторий на обжалование не применяется. "
            f"При 1 участнике (без аукциона) акцепт и подписание возможны {dt_end.strftime('%d.%m.%Y')} – {next_bd_str}. "
            f"При 2 и более участниках аукцион пройдет в ближайший рабочий день ({next_bd_str}), "
            f"а подписание договора возможно сразу после акцепта ({next_bd_str} или {next_business_day(next_bd, 1).strftime('%d.%m.%Y')})."
        )
    elif method_type == "aboveThreshold":
        moratorium_from_d = (dt_end + timedelta(days=6)).strftime("%d.%m.%Y")
        moratorium_from_auc = (next_bd + timedelta(days=6)).strftime("%d.%m.%Y")
        return (
            f"Процедура: Открытые торги с особенностями (Постановление КМУ №1178). "
            f"Окончание подачи предложений (День D): {d_str}. "
            f"Сценарий А (если подано 2 и более предложений — с аукционом): аукцион назначается на рабочий день ({next_bd_str}). "
            f"При публикации уведомления о намерении заключить договор в день аукциона ({next_bd_str}) и мгновенной загрузке документов победителя (п. 47 Особливостей), "
            f"действует императивный минимальный срок п. 49 Постановления КМУ №1178 (не ранее чем через 5 календарных дней с даты публикации уведомления о намерении, ст. 253 ГКУ): "
            f"ближайшая законная дата заключения договора — {moratorium_from_auc}. "
            f"Сценарий Б (если мы единственный участник — 1 предложение без аукциона по п. 36 Особливостей): "
            f"раскрытие сразу {dt_end.strftime('%d.%m.%Y')}, ближайшая дата подписания — {next_bd_str} (либо {moratorium_from_d}, если электронная площадка включает 5-дневный блок п. 49)."
        )
    else:
        moratorium_10d = (next_bd + timedelta(days=11)).strftime("%d.%m.%Y")
        return (
            f"Процедура: Открытые торги (ст. 33 Закона Украины 'О публичных закупках'). "
            f"Окончание подачи: {d_str}. С учетом 10-дневного срока на обжалование со дня уведомления о намерении заключить договор, "
            f"ближайшая законная дата заключения договора — {moratorium_10d}."
        )

def parse_docx_bytes(file_bytes: bytes, filename: str) -> str:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        lines = [f"=== ДОКУМЕНТ: {filename} ==="]
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                lines.append(txt)
        for table in doc.tables:
            lines.append("--- Таблица ---")
            for row in table.rows:
                cells = []
                last_txt = None
                for cell in row.cells:
                    c_txt = " ".join(cell.text.split())
                    if c_txt and c_txt != last_txt:
                        cells.append(c_txt)
                        last_txt = c_txt
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)[:75000]
    except Exception as e:
        return f"Ошибка чтения DOCX {filename}: {e}"

def parse_old_doc_bytes(file_bytes: bytes, filename: str) -> str:
    try:
        with tempfile.NamedTemporaryFile(suffix=".doc", delete=True) as tmp:
            tmp.write(file_bytes)
            tmp.flush()
            res = subprocess.run(["antiword", tmp.name], capture_output=True, text=True, timeout=10)
            if res.returncode == 0 and res.stdout.strip():
                return f"=== ДОКУМЕНТ (DOC): {filename} ===\n" + res.stdout[:75000]
    except Exception:
        pass
    return ""

def extract_all_attachments(documents_json: list) -> tuple[str, list]:
    latest_docs = {}
    for d in documents_json:
        if d.get("documentOf") not in ("tender", "lot", "item"):
            continue
        doc_id = d.get("id")
        if not doc_id:
            continue
        if doc_id not in latest_docs or d.get("dateModified", "") > latest_docs[doc_id].get("dateModified", ""):
            latest_docs[doc_id] = d

    valid_docs = []
    for d in latest_docs.values():
        title = d.get("title", "").lower()
        if title.endswith((".p7s", ".asic", ".sig", ".asice")):
            continue
        valid_docs.append(d)

    def doc_priority(doc_item):
        t = doc_item.get("title", "").lower()
        if any(k in t for k in ("догов", "контракт", "техніч", "документац", "тд", "додаток", "специфікац", "вимог", "оголош")):
            return 0
        return 1

    valid_docs.sort(key=doc_priority)

    extracted_texts = []
    pdf_parts = []

    for d in valid_docs[:8]:
        title = d.get("title", "document")
        url = d.get("url", "")
        if not url:
            continue

        try:
            resp = http_session.get(url, timeout=25)
            if resp.status_code != 200:
                continue
            content = resp.content
            low_title = title.lower()

            if low_title.endswith(".docx"):
                extracted_texts.append(parse_docx_bytes(content, title))
            elif low_title.endswith(".doc"):
                doc_txt = parse_old_doc_bytes(content, title)
                if doc_txt:
                    extracted_texts.append(doc_txt)
            elif low_title.endswith(".pdf"):
                if len(content) < 15 * 1024 * 1024 and len(pdf_parts) < 3:
                    pdf_parts.append(types.Part.from_bytes(data=content, mime_type="application/pdf"))
                    extracted_texts.append(f"[Прикреплен PDF-документ: {title}]")
            elif low_title.endswith(".zip"):
                with zipfile.ZipFile(io.BytesIO(content)) as zf:
                    for inner_name in zf.namelist()[:8]:
                        low_inner = inner_name.lower()
                        if low_inner.endswith((".p7s", ".sig")):
                            continue
                        inner_bytes = zf.read(inner_name)
                        if low_inner.endswith(".docx"):
                            extracted_texts.append(parse_docx_bytes(inner_bytes, f"{title}/{inner_name}"))
                        elif low_inner.endswith(".doc"):
                            doc_txt = parse_old_doc_bytes(inner_bytes, f"{title}/{inner_name}")
                            if doc_txt:
                                extracted_texts.append(doc_txt)
                        elif low_inner.endswith(".pdf") and len(inner_bytes) < 15 * 1024 * 1024 and len(pdf_parts) < 3:
                            pdf_parts.append(types.Part.from_bytes(data=inner_bytes, mime_type="application/pdf"))
                            extracted_texts.append(f"[Прикреплен PDF из архива: {inner_name}]")
        except Exception as e:
            extracted_texts.append(f"Ошибка загрузки файла {title}: {e}")

    return "\n\n".join(extracted_texts), pdf_parts

# Когнитивная экспертиза нового тендера через Gemini
def analyze_tender_with_gemini(tender_data: dict, docs_text: str, pdf_parts: list) -> dict:
    val = tender_data.get("value", {})
    budget_amount = val.get("amount", 0)
    vat_included_flag = val.get("valueAddedTaxIncluded", False)
    guarantee_amount = tender_data.get("guarantee", {}).get("amount", 0)
    method_type = tender_data.get("procurementMethodType", "")
    method_readable = PROCUREMENT_METHODS.get(method_type, "Открытые торги")
    statutory_dates_hint = build_statutory_dates_reference(tender_data)

    # Для процедуры aboveThreshold («Відкриті торги з особливостями») в объявлении Prozorro параметр
    # «Ціна пропозиції обмежена очікуваною вартістю» всегда равен «Ні»
    is_price_limited_by_announcement = "Ні" if method_type == "aboveThreshold" else "Так"

    items_json = json.dumps(tender_data.get("items", []), ensure_ascii=False)
    lots_json = json.dumps(tender_data.get("lots", []), ensure_ascii=False)
    e_terms_json = json.dumps(tender_data.get("milestones", []), ensure_ascii=False)

    prompt = f"""
    ПАСПОРТ ЗАКУПКИ ИЗ PROZORRO:
    ID закупки: {tender_data.get('tenderID', '')}
    Заголовок: {tender_data.get('title', '')}
    Описание: {tender_data.get('description', '')}
    Вид процедуры: {method_readable} ({method_type})
    Параметр объявления «Ціна пропозиції обмежена очікуваною вартістю»: {is_price_limited_by_announcement}
    Бюджет закупки в системе Prozorro: {budget_amount} UAH (флаг valueAddedTaxIncluded: {vat_included_flag})
    Тендерное обеспечение в паспорте Prozorro: {guarantee_amount} UAH
    Справка по законным срокам процедуры: {statutory_dates_hint}
    Лоты объявления: {lots_json}
    Номенклатура объявления: {items_json}
    Электронные условия объявления Prozorro: {e_terms_json}

    ПОЛНЫЙ ТЕКСТ ТЕНДЕРНОЙ ДОКУМЕНТАЦИИ И ПРОЕКТА ДОГОВОРА:
    {docs_text[:130000]}
    """

    contents = [prompt] + pdf_parts
    gen_config = types.GenerateContentConfig(
        system_instruction=(
            "Ты — ведущий тендерный юрист и аналитик оптового поставщика нефтепродуктов в Украине. "
            "Проанализируй паспорт закупки и прикрепленный проект договора/ТД.\n"
            "КРИТЕРИЙ ОТБОРА (is_relevant = true):\n"
            "1. Ищем ТОЛЬКО закупки БЕНЗИНА или ДИЗЕЛЬНОГО ТОПЛИВА НАЛИВОМ: "
            "автомобильным транспортом (бензовозами/автоцистернами или самовывозом) ИЛИ железнодорожным транспортом (вагон-цистернами).\n"
            "2. Если закупается авиатопливо (JET A-1, РТ, ТС-1), керосин (гас), газ, сжиженный газ, мазут или масла — СТРОГО верни is_relevant = false!\n"
            "3. Если основной предмет — талоны, топливные карты, скретч-карты или розничная заправка на АЗС — СТРОГО верни is_relevant = false!\n"
            "4. ПРОВЕРКА НДС (vat_rate_percent): Проверь проект договора и условия закупки. Если применяется ставка НДС 0% (ПДВ 0%, в т.ч. по Постановлению КМУ №178) — укажи vat_rate_percent = 0. Если стандартный НДС 20% — укажи vat_rate_percent = 20.\n"
            "5. ПРОВЕРКА ПРЕВЫШЕНИЯ ОЖИДАЕМОЙ СТОИМОСТИ (over_budget_policy и over_budget_percent):\n"
            "   Изучи Тендерную документацию (особенно Раздел III 'Інструкція з підготовки тендерної пропозиції', Раздел V 'Оцінка тендерної пропозиції' и Додаток 'Методика оцінки') на предмет условия 'Прийнятний відсоток перевищення ціни тендерної пропозиції учасника над очікуваною вартістю предмета закупівлі' (п. 28 / п. 44 Постановления КМУ №1178):\n"
            "   - Если в ТД указан конкретный допустимый процент превышения ожидаемой стоимости (например, 'не більше ніж 10%'): установи over_budget_policy = 'ALLOWED_WITH_PERCENT' и over_budget_percent = число процента (например, 10).\n"
            "   - Если в объявлении стоит 'Ціна пропозиції обмежена очікуваною вартістю: Ні' (открытые торги с особенностями), но в тексте ТД заказчик ПРЯМО ЗАПРЕТИЛ подавать предложения выше ожидаемой стоимости (например, написано 'Замовник не приймає до розгляду тендерну пропозицію, ціна якої є вищою, ніж очікувана вартість' или 'ціна пропозиції не може перевищувати очікувану вартість'): установи over_budget_policy = 'PROHIBITED_IN_TD' и over_budget_percent = 0.\n"
            "   - Если в объявлении стоит 'Ні', но ни в объявлении, ни в ТД процент превышения не указан и прямого запрета нет: установи over_budget_policy = 'UNSPECIFIED_IN_TD' и over_budget_percent = 0.\n"
            "   - Если в объявлении стоит 'Так' (цена ограничена ожидаемой стоимостью): установи over_budget_policy = 'LIMITED_BY_ANNOUNCEMENT' и over_budget_percent = 0.\n"
            "6. ЗАПРЕТ НА КОД ДК В ТЕКСТЕ ПУНКТА 1: В поле subject_grade_and_volume НЕ ПИШИ код классификатора ДК 021:2015 (например, не пиши 'код ДК 021:2015:09130000-9'). Указывай только наименование топлива, марку, сезонность и объем.\n"
            "7. ЗАПРЕТ НА АНГЛИЦИЗМЫ: Никогда не используй английские названия полей Prozorro ('Milestones', 'Items', 'Lots'). Вместо 'п. 1 Milestones' пиши 'п. 1 Электронных условий объявления', вместо 'Items' — 'Номенклатура объявления'.\n"
            "8. ФОРМУЛА ЦЕНЫ: Если цена формульная, выдели отдельно строку предварительной формулы (если она есть), строку окончательной формулы (одной строкой без переносов) и массив расшифровки каждого значения формулы.\n"
            "9. РАСЧЕТ РАННЕЙ ДАТЫ ДОГОВОРА: Рассчитай самую раннюю законную дату подписания договора при нашей победе и максимальной оперативности заказчика. В случае противоречия сроков между ТД и законодательством приоритет у законодательства Украины.\n"
            "Все ответы формируй строго на русском языке без эмодзи."
        ),
        response_mime_type="application/json",
        response_schema={
            "type": "OBJECT",
            "properties": {
                "is_relevant": {"type": "BOOLEAN"},
                "vat_rate_percent": {"type": "INTEGER"},
                "over_budget_policy": {
                    "type": "STRING",
                    "description": "Одно из значений: 'ALLOWED_WITH_PERCENT', 'PROHIBITED_IN_TD', 'UNSPECIFIED_IN_TD', 'LIMITED_BY_ANNOUNCEMENT'"
                },
                "over_budget_percent": {
                    "type": "NUMBER",
                    "description": "Процент допустимого превышения ожидаемой стоимости по ТД (например, 10), либо 0"
                },
                "subject_grade_and_volume": {
                    "type": "STRING",
                    "description": "Пункт 1: Наименование товара, марка топлива и точное количество с единицами измерения и ссылкой на пункт Договора/Объявления БЕЗ указания кода ДК 021:2015"
                },
                "payment_terms": {"type": "STRING"},
                "is_formula_pricing": {"type": "BOOLEAN"},
                "preliminary_formula_line": {
                    "type": "STRING",
                    "description": "Формула предварительной цены одной строкой (если есть в договоре), иначе пустая строка ''"
                },
                "final_formula_line": {
                    "type": "STRING",
                    "description": "Полная формула окончательной цены одной строкой с указанием пункта Договора (или 'Цена фиксированная')"
                },
                "formula_variables": {
                    "type": "ARRAY",
                    "description": "Расшифровка каждой переменной формулы по отдельности",
                    "items": {
                        "type": "OBJECT",
                        "properties": {
                            "symbol": {"type": "STRING", "description": "Обозначение переменной в формуле, например: P, Platts, F, Кдолар, Гк"},
                            "explanation": {"type": "STRING", "description": "Точная расшифровка значения этой переменной по договору"}
                        },
                        "required": ["symbol", "explanation"]
                    }
                },
                "delivery_mode_and_location": {"type": "STRING"},
                "delivery_and_contract_term": {"type": "STRING"},
                "bid_security_uah": {"type": "STRING"},
                "contract_security_details": {"type": "STRING"},
                "earliest_contract_date_calculation": {
                    "type": "STRING",
                    "description": "Подробный юридический расчет самой ранней даты заключения договора по этапам закупки (с учетом вида процедуры, наличия/отсутствия аукциона, моратория на обжалование и обеспечения договора)"
                },
                "features_pitfalls_and_risks": {
                    "type": "STRING",
                    "description": "Раздел 1 Особенностей тендера: Особенности приема-передачи товара, риски поставщика (с номерами пунктов Договора/ТД и цитатами)"
                },
                "features_penalties": {
                    "type": "STRING",
                    "description": "Раздел 2 Особенностей тендера: Ответственность поставщика (штрафные санкции и пеня) с номерами пунктов"
                },
                "features_price_increase_rules": {
                    "type": "STRING",
                    "description": "Раздел 3 Особенностей тендера: Условия и ограничения для повышения цены с номерами пунктов"
                }
            },
            "required": [
                "is_relevant", "vat_rate_percent", "over_budget_policy", "over_budget_percent",
                "subject_grade_and_volume", "payment_terms", "is_formula_pricing",
                "preliminary_formula_line", "final_formula_line", "formula_variables",
                "delivery_mode_and_location", "delivery_and_contract_term",
                "bid_security_uah", "contract_security_details",
                "earliest_contract_date_calculation",
                "features_pitfalls_and_risks", "features_penalties", "features_price_increase_rules"
            ]
        },
        temperature=0.0
    )

    last_err = None
    for loc_name, v_client in vertex_clients:
        for model_name in GEMINI_MODELS_FALLBACK:
            for attempt in range(2):
                try:
                    response = v_client.models.generate_content(
                        model=model_name,
                        contents=contents,
                        config=gen_config
                    )
                    return json.loads(response.text)
                except Exception as err:
                    last_err = err
                    err_str = str(err)
                    if "404" in err_str or "NOT_FOUND" in err_str:
                        break
                    if "503" in err_str or "UNAVAILABLE" in err_str or "429" in err_str:
                        time.sleep(3 * (attempt + 1))
                        continue
                    break

    raise (last_err or Exception("Не удалось получить ответ от моделей Gemini"))

# Краткое резюме причины отклонения участника через Gemini
def summarize_rejection_reason_with_gemini(tender_ua_id: str, supplier_name: str, award_obj: dict) -> str:
    raw_title = str(award_obj.get("title", "") or "")
    raw_desc = str(award_obj.get("description", "") or "")
    docs = award_obj.get("documents", []) or []

    extracted_txt, pdf_parts = extract_all_attachments(docs[:3]) if docs else ("", [])
    if not extracted_txt and not pdf_parts and len(raw_desc) > 15 and "протокол" not in raw_desc.lower():
        return clean_and_enforce_source(raw_desc, "Протокола отклонения")

    prompt = (
        f"Тендер: {tender_ua_id}. Участник '{supplier_name}' был отклонен заказчиком.\n"
        f"Заголовок в системе: {raw_title}\nОписание в системе: {raw_desc}\n"
        f"Текст протокола отклонения:\n{extracted_txt[:30000]}\n\n"
        f"Сформулируй в 1–2 емких предложениях на русском языке точную юридическую и фактическую причину отклонения этого участника заказчиком."
    )
    contents = [prompt] + pdf_parts
    for _, v_client in vertex_clients:
        for model_name in GEMINI_MODELS_FALLBACK:
            try:
                resp = v_client.models.generate_content(
                    model=model_name,
                    contents=contents,
                    config=types.GenerateContentConfig(temperature=0.0)
                )
                if resp and resp.text:
                    return resp.text.strip()
            except Exception:
                continue
    fallback = f"{raw_title} {raw_desc}".strip()
    return fallback if fallback else "Причина указана в прикрепленном протоколе рассмотрения в Prozorro."

# Форматирование текста формулы цены (формула одной строкой + каждая переменная с новой строки жирным)
def build_formatted_formula_html(ai: dict) -> str:
    if not ai.get("is_formula_pricing"):
        return "Цена фиксированная."

    prelim = clean_and_enforce_source(ai.get("preliminary_formula_line", ""), "Договора")
    final_f = clean_and_enforce_source(ai.get("final_formula_line", ""), "Договора")
    variables = ai.get("formula_variables", []) or []

    blocks = []
    if prelim and prelim.lower() not in ("нет", "не предусмотрена", "отсутствует"):
        blocks.append(f"<b>Формула предварительной цены:</b>\n<code>{html.escape(prelim)}</code>")

    if final_f:
        blocks.append(f"<b>Формула окончательной цены:</b>\n<code>{html.escape(final_f)}</code>")

    if variables:
        var_lines = ["<b>Расшифровка значений формулы:</b>"]
        for v in variables:
            sym = str(v.get("symbol", "")).strip()
            exp = clean_and_enforce_source(str(v.get("explanation", "")).strip(), "Договора")
            if sym and exp:
                var_lines.append(f"• <b>{html.escape(sym)}</b> — {html.escape(exp)}")
        blocks.append("\n".join(var_lines))

    return "\n\n".join(blocks) if blocks else html.escape(final_f or "Формула указана в проекте договора.")

# Сборка карточки нового тендера (строгий деловой текст без иконок внутри карточки)
def dispatch_to_crm(tender_data: dict, ai: dict, docs_text: str):
    tender_ua_id = tender_data.get("tenderID", "Не указан")
    entity = tender_data.get("procuringEntity", {})
    cust_name = entity.get("name", "Не указан")
    cust_edrpou = str(entity.get("identifier", {}).get("id", "Не указан"))
    customer_str = f"{cust_name} (ЕГРПОУ: {cust_edrpou})"

    method_type = tender_data.get("procurementMethodType", "")
    procedure_clean = PROCUREMENT_METHODS.get(method_type, "Открытые торги")

    pricing_meta = extract_pricing_meta_and_budget(tender_data, ai, docs_text)
    budget_display_str = pricing_meta["budgetDisplay"]

    raw_end_date = tender_data.get("tenderPeriod", {}).get("endDate", "")
    deadline_formatted = format_deadline_strict(raw_end_date)
    tender_url = build_smarttender_url(tender_ua_id)

    subject_vol_clean = clean_and_enforce_source(ai.get("subject_grade_and_volume", ""), "Договора", strip_cpv=True)
    payment_terms_clean = clean_and_enforce_source(ai.get("payment_terms", ""), "Договора")
    delivery_loc_clean = clean_and_enforce_source(ai.get("delivery_mode_and_location", ""), "Договора")
    delivery_term_clean = clean_and_enforce_source(ai.get("delivery_and_contract_term", ""), "Договора")
    bid_sec_clean = clean_and_enforce_source(ai.get("bid_security_uah", ""), "Тендерной документации")
    contract_sec_clean = clean_and_enforce_source(ai.get("contract_security_details", ""), "Договора")

    pitfalls_clean = clean_and_enforce_source(ai.get("features_pitfalls_and_risks", ""), "Договора")
    penalties_clean = clean_and_enforce_source(ai.get("features_penalties", ""), "Договора")
    price_inc_clean = clean_and_enforce_source(ai.get("features_price_increase_rules", ""), "Договора")
    earliest_date_calc = clean_and_enforce_source(ai.get("earliest_contract_date_calculation", ""), "Законодательства и ТД")

    is_formula = bool(ai.get("is_formula_pricing"))
    if is_formula:
        pricing_html = "Цена формульная (смотри дополнительно по кнопке «Формула цены» под данной карточкой тендера)"
    else:
        pricing_html = "Цена фиксированная"

    formatted_formula_html = build_formatted_formula_html(ai)

    # Строгий деловой текст карточки нового тендера (0 иконок, без кода ДК в п.1, с расчетом превышения бюджета в п.2)
    card_html = (
        f"<b>Тендер:</b> <code>{html.escape(tender_ua_id)}</code> ({html.escape(procedure_clean)})\n"
        f"-----------------------------------------\n"
        f"<b>{html.escape(customer_str)}</b>\n\n"
        f"<b>1) Предмет закупки и объем:</b> {html.escape(subject_vol_clean)}\n\n"
        f"<b>2) Бюджет закупки:</b> {html.escape(budget_display_str)}\n\n"
        f"<b>3) Условия оплаты:</b> {html.escape(payment_terms_clean)}\n\n"
        f"<b>4) Ценообразование:</b> {pricing_html}\n\n"
        f"<b>5) Срок подачи тендерного предложения:</b> до {html.escape(deadline_formatted)}\n\n"
        f"<b>6) Условия поставки и адрес передачи товара:</b> {html.escape(delivery_loc_clean)}\n\n"
        f"<b>7) Срок поставки и срок действия договора:</b> {html.escape(delivery_term_clean)}\n\n"
        f"<b>8) Обеспечение тендерного предложения:</b> {html.escape(bid_sec_clean)}\n\n"
        f"<b>9) Обеспечение договора:</b> {html.escape(contract_sec_clean)}\n\n"
        f"<b>10) Ссылка на тендер:</b> <a href=\"{html.escape(tender_url)}\">{html.escape(tender_url)}</a>"
    )

    features_detailed_text = (
        f"1. ОСОБЕННОСТИ ПРИЕМА-ПЕРЕДАЧИ ТОВАРА, РИСКИ ПОСТАВЩИКА:\n"
        f"{pitfalls_clean}\n\n"
        f"2. ОТВЕТСТВЕННОСТЬ ПОСТАВЩИКА (ШТРАФНЫЕ САНКЦИИ И ПЕНЯ):\n"
        f"{penalties_clean}\n\n"
        f"3. УСЛОВИЯ И ОГРАНИЧЕНИЯ ДЛЯ ПОВЫШЕНИЯ ЦЕНЫ (ДОКУМЕНТЫ И ПУНКТЫ ДОГОВОРА):\n"
        f"{price_inc_clean}"
    )

    doc_snapshot = {
        d.get("id"): f"{d.get('title', '')}|{d.get('dateModified', '')}"
        for d in (tender_data.get("documents") or [])
        if d.get("id") and not str(d.get("title", "")).lower().endswith((".p7s", ".sig", ".asic", ".asice"))
    }

    payload = {
        "source": "python_agent",
        "secret": SECRET_BRIDGE_KEY,
        "tender": {
            "prozorroInternalId": tender_data.get("id", tender_ua_id),
            "tenderID": tender_ua_id,
            "customerName": cust_name,
            "customerEdrpou": cust_edrpou,
            "procedureType": procedure_clean,
            "budgetDisplay": budget_display_str,
            "isZeroVat": pricing_meta["isZeroVat"],
            "vatRatePercent": pricing_meta["vatRatePercent"],
            "itemsCount": pricing_meta["itemsCount"],
            "singleItemQty": pricing_meta["singleItemQty"],
            "unitLabel": pricing_meta["unitLabel"],
            "fuelSummary": subject_vol_clean,
            "deliveryAndTerms": f"{delivery_loc_clean} | Сроки: {delivery_term_clean}",
            "paymentAndGuarantees": (
                f"Оплата: {payment_terms_clean} | "
                f"Тендерная гарантия: {bid_sec_clean} | "
                f"Обеспечение договора: {contract_sec_clean}"
            ),
            "deadline": deadline_formatted,
            "deadlineIso": raw_end_date,
            "tenderUrl": tender_url,
            "isFormula": is_formula,
            "formulaText": formatted_formula_html,
            "featuresDetailedText": features_detailed_text,
            "earliestDateText": earliest_date_calc,
            "docSnapshot": doc_snapshot,
            "rawBudgetAmount": float(tender_data.get("value", {}).get("amount", 0) or 0),
            "telegramCardHtml": card_html
        }
    }

    resp = http_session.post(GOOGLE_WEB_APP_URL, json=payload, timeout=25)
    print(f" -> Ответ Google Таблицы: {resp.text[:120]}")

# Формирование списка участников со статусными бизнес-иконками
def extract_participants_from_prozorro(
    t_data: dict,
    is_zero_vat: bool,
    items_count: int,
    single_qty: float,
    unit_label: str
) -> tuple[list, int]:
    awards = t_data.get("awards", []) or []
    bids = t_data.get("bids", []) or []

    if items_count <= 0:
        items = t_data.get("items", [])
        fuel_items = [
            i for i in items
            if str(i.get("classification", {}).get("id", "")).startswith(TARGET_CPV_PREFIXES)
            and not str(i.get("classification", {}).get("id", "")).startswith(EXCLUDED_CPV_PREFIXES)
        ]
        target_items = fuel_items if fuel_items else items
        items_count = len(target_items)
        if items_count == 1:
            single_qty = float(target_items[0].get("quantity", 0) or 0)
            unit_raw = str(target_items[0].get("unit", {}).get("name", "ед.")).lower()
            if "літр" in unit_raw or "литр" in unit_raw or unit_raw == "л":
                unit_label = "литр"
            elif "тон" in unit_raw or unit_raw == "т":
                unit_label = "тонну"
            elif "кілогр" in unit_raw or "килогр" in unit_raw or unit_raw == "кг":
                unit_label = "кг"
            else:
                unit_label = unit_raw

    award_by_bid = {}
    award_by_edrpou = {}
    for aw in awards:
        aw_status = aw.get("status", "")
        bid_id = aw.get("bid_id", "")
        supp = (aw.get("suppliers") or [{}])[0]
        edrpou = str(supp.get("identifier", {}).get("id", ""))
        has_24h = any(m.get("code") == "24h" for m in (aw.get("milestones") or []))

        label = "▫️ Участник"
        if aw_status == "active":
            label = "🏆 Признан победителем (ожидается заключение договора)"
        elif aw_status == "pending":
            if has_24h:
                label = "⏱ На рассмотрении (дано 24 часа на исправление)"
            else:
                label = "⏳ На рассмотрении (оценка предложения заказчиком)"
        elif aw_status == "unsuccessful":
            label = "❌ Предложение отклонено"
        elif aw_status == "cancelled":
            label = "▫️ Решение отменено"

        if bid_id:
            award_by_bid[bid_id] = label
        if edrpou:
            award_by_edrpou[edrpou] = label

    participants = []
    if bids:
        for idx, b in enumerate(bids, start=1):
            if b.get("status", "active") not in ("active", "pending"):
                continue
            tenderers = b.get("tenderers") or []
            if tenderers:
                supp = tenderers[0]
                name = supp.get("name", f"Участник №{idx}")
                edrpou = str(supp.get("identifier", {}).get("id", "скрыт до конца аукциона"))
            else:
                name = f"Участник №{idx} (наименование скрыто Prozorro до окончания аукциона)"
                edrpou = "скрыт"

            val = b.get("value")
            if not val and b.get("lotValues"):
                val = b["lotValues"][0].get("value", {})

            if val and val.get("amount") is not None:
                raw_amt = float(val.get("amount", 0) or 0)
                vat_inc = bool(val.get("valueAddedTaxIncluded", False))
                price_formatted = format_participant_price_with_vat(
                    raw_amt, vat_inc, is_zero_vat, items_count, single_qty, unit_label
                )
            else:
                raw_amt = 999999999999.0 + idx
                price_formatted = "Стартовая цена скрыта Prozorro до завершения аукциона"

            bid_id = b.get("id", "")
            status_label = award_by_bid.get(bid_id) or award_by_edrpou.get(edrpou) or "▫️ Участник"

            participants.append({
                "rawAmount": raw_amt,
                "name": name,
                "edrpou": edrpou,
                "priceFormatted": price_formatted,
                "statusLabel": status_label
            })

    if not participants and awards:
        seen_edrpou = set()
        for aw in awards:
            aw_status = aw.get("status", "")
            if aw_status == "cancelled":
                continue
            supp = (aw.get("suppliers") or [{}])[0]
            name = supp.get("name", "Не указано")
            edrpou = str(supp.get("identifier", {}).get("id", "Не указан"))
            val = aw.get("value", {})
            if val.get("amount") is None:
                continue
            raw_amt = float(val.get("amount", 0) or 0)
            vat_inc = bool(val.get("valueAddedTaxIncluded", False))
            has_24h = any(m.get("code") == "24h" for m in (aw.get("milestones") or []))

            if aw_status == "active":
                status_label = "🏆 Признан победителем (ожидается заключение договора)"
            elif aw_status == "pending":
                status_label = "⏱ На рассмотрении (дано 24 часа на исправление)" if has_24h else "⏳ На рассмотрении (оценка предложения заказчиком)"
            elif aw_status == "unsuccessful":
                status_label = "❌ Предложение отклонено"
            else:
                status_label = "▫️ Участник"

            key = f"{edrpou}_{raw_amt}"
            if key in seen_edrpou:
                continue
            seen_edrpou.add(key)

            price_formatted = format_participant_price_with_vat(
                raw_amt, vat_inc, is_zero_vat, items_count, single_qty, unit_label
            )
            participants.append({
                "rawAmount": raw_amt,
                "name": name,
                "edrpou": edrpou,
                "priceFormatted": price_formatted,
                "statusLabel": status_label
            })

    participants.sort(key=lambda x: x["rawAmount"])
    return participants, items_count

def fetch_dasu_events_for_tenders(tracked_internal_ids: set) -> dict:
    dasu_by_tender = {}
    if not tracked_internal_ids:
        return dasu_by_tender
    try:
        r = http_session.get(
            f"{PROZORRO_API_BASE}/monitorings?descending=1&limit=100&opt_fields=tender_id,status",
            timeout=10
        )
        if r.status_code != 200:
            return dasu_by_tender
        items = r.json().get("data", [])
        for m in items:
            t_int_id = m.get("tender_id")
            if t_int_id in tracked_internal_ids:
                m_id = m.get("id")
                mr = http_session.get(f"{PROZORRO_API_BASE}/monitorings/{m_id}", timeout=10)
                if mr.status_code == 200:
                    dasu_by_tender[t_int_id] = mr.json().get("data", {})
    except Exception:
        pass
    return dasu_by_tender

# Комплексное сопровождение тендеров из списка «Принимаем участие» + фоновое обогащение истории базы
def check_tracked_and_background_tenders():
    try:
        req_payload = {
            "source": "python_agent",
            "secret": SECRET_BRIDGE_KEY,
            "action": "get_tracked_tenders"
        }
        r = http_session.post(GOOGLE_WEB_APP_URL, json=req_payload, timeout=20)
        if r.status_code != 200:
            return
        resp_data = r.json()
        tracked_list = resp_data.get("trackedTenders", []) or []
        bg_list = resp_data.get("backgroundTenders", []) or []

        now_utc = datetime.now(timezone.utc)
        tracked_internal_ids = {item.get("prozorroInternalId") for item in tracked_list if item.get("prozorroInternalId")}
        dasu_map = fetch_dasu_events_for_tenders(tracked_internal_ids)

        if tracked_list:
            print(f"=== Сопровождение тендеров «Принимаем участие» ({len(tracked_list)} шт.) ===")

        for item in tracked_list:
            tender_ua_id = item.get("tenderID", "")
            lookup_id = item.get("prozorroInternalId") or tender_ua_id
            is_zero_vat = bool(item.get("isZeroVat", False))
            items_count = int(item.get("itemsCount", 0) or 0)
            single_qty = float(item.get("singleItemQty", 0) or 0)
            unit_label = str(item.get("unitLabel", "литр") or "литр")
            state = item.get("trackingState", {}) or {}

            t_data = fetch_tender_detail(lookup_id)
            if not t_data:
                continue

            t_status = t_data.get("status", "")
            awards = t_data.get("awards", []) or []
            contracts = t_data.get("contracts", []) or []
            participants, items_count = extract_participants_from_prozorro(
                t_data, is_zero_vat, items_count, single_qty, unit_label
            )

            events_for_all = []
            events_for_admin = []
            state_changed = False

            # 1. Контроль изменений в тендере
            cur_docs = {
                d.get("id"): f"{d.get('title', '')}|{d.get('dateModified', '')}"
                for d in (t_data.get("documents") or [])
                if d.get("id") and d.get("documentOf") in ("tender", "lot", "item")
                and not str(d.get("title", "")).lower().endswith((".p7s", ".sig", ".asic", ".asice"))
            }
            old_docs = state.get("docSnapshot") or {}
            new_end_iso = t_data.get("tenderPeriod", {}).get("endDate", "")
            old_end_iso = state.get("deadlineIso", "")
            new_budget_raw = float(t_data.get("value", {}).get("amount", 0) or 0)
            old_budget_raw = float(state.get("rawBudgetAmount", 0) or 0)

            changes_desc = []
            if old_docs:
                for d_id, d_val in cur_docs.items():
                    title_part = d_val.split("|")[0]
                    if d_id not in old_docs:
                        changes_desc.append(f"• Загружен новый файл: <b>{html.escape(title_part)}</b>")
                    elif old_docs[d_id] != d_val:
                        changes_desc.append(f"• Обновлен файл: <b>{html.escape(title_part)}</b>")
            if old_end_iso and new_end_iso and old_end_iso != new_end_iso:
                changes_desc.append(f"• Изменен срок подачи предложений: <b>до {html.escape(format_deadline_strict(new_end_iso))}</b>")
            if old_budget_raw > 0 and new_budget_raw > 0 and abs(new_budget_raw - old_budget_raw) > 1:
                changes_desc.append(f"• Изменен бюджет закупки: <b>{new_budget_raw:,.2f} грн</b>".replace(",", " "))

            if changes_desc:
                events_for_all.append({
                    "title": "📝 ЗАКАЗЧИК ВНЕС ИЗМЕНЕНИЯ В ТЕНДЕР",
                    "bodyHtml": "\n".join(changes_desc),
                    "newDeadlineFormatted": format_deadline_strict(new_end_iso) if new_end_iso else ""
                })
                state["docSnapshot"] = cur_docs
                state["deadlineIso"] = new_end_iso
                state["rawBudgetAmount"] = new_budget_raw
                state_changed = True
            elif not old_docs:
                state["docSnapshot"] = cur_docs
                state["deadlineIso"] = new_end_iso
                state["rawBudgetAmount"] = new_budget_raw
                state_changed = True

            # 2. Напоминание Админу за 12 часов до окончания подачи предложений
            dt_end = parse_iso_dt(new_end_iso)
            if dt_end and t_status in ("active.enquiries", "active.tendering"):
                hours_left = (dt_end - now_utc).total_seconds() / 3600.0
                rem_key = f"rem12h_{new_end_iso}"
                if 0 < hours_left <= 12.0 and state.get("notified12hKey") != rem_key:
                    kyiv_dt = dt_end.astimezone(timezone(timedelta(hours=2)))
                    day_word = "Сегодня" if kyiv_dt.date() == now_utc.astimezone(timezone(timedelta(hours=2))).date() else "Завтра"
                    events_for_admin.append({
                        "title": "⏳ НАПОМИНАНИЕ О ДЕДЛАЙНЕ ПОДАЧИ (МЕНЕЕ 12 ЧАСОВ)",
                        "bodyHtml": f"<b>{day_word} в {kyiv_dt.strftime('%H:%M')} ({kyiv_dt.strftime('%d-%m-%Y')})</b> закрывается подача предложений по данному тендеру."
                    })
                    state["notified12hKey"] = rem_key
                    state_changed = True

            # 3. Уведомление Админу о назначении времени аукциона и за 1 час до начала аукциона
            auc_start_iso = t_data.get("auctionPeriod", {}).get("startDate", "")
            if not auc_start_iso and t_data.get("lots"):
                auc_start_iso = (t_data["lots"][0].get("auctionPeriod") or {}).get("startDate", "")
            auc_url = t_data.get("auctionUrl", "") or build_smarttender_url(tender_ua_id)

            if auc_start_iso:
                if state.get("notifiedAuctionSet") != auc_start_iso:
                    events_for_admin.append({
                        "title": "🔔 НАЗНАЧЕНО ВРЕМЯ ЭЛЕКТРОННОГО АУКЦИОНА",
                        "bodyHtml": (
                            f"• <b>Дата и время старта аукциона:</b> <b>{html.escape(format_deadline_strict(auc_start_iso))}</b>\n"
                            f"• <b>Ссылка на вход в аукцион:</b> <a href=\"{html.escape(auc_url)}\">{html.escape(auc_url)}</a>"
                        )
                    })
                    state["notifiedAuctionSet"] = auc_start_iso
                    state_changed = True

                dt_auc = parse_iso_dt(auc_start_iso)
                if dt_auc:
                    mins_to_auc = (dt_auc - now_utc).total_seconds() / 60.0
                    auc_1h_key = f"auc1h_{auc_start_iso}"
                    if 0 < mins_to_auc <= 65.0 and state.get("notifiedAuction1h") != auc_1h_key:
                        events_for_admin.append({
                            "title": "🔔 НАПОМИНАНИЕ: ЧЕРЕЗ 1 ЧАС СТАРТ АУКЦИОНА",
                            "bodyHtml": (
                                f"• <b>Старт аукциона:</b> <b>{html.escape(format_deadline_strict(auc_start_iso))}</b>\n"
                                f"• <b>Ссылка на вход в аукцион:</b> <a href=\"{html.escape(auc_url)}\">{html.escape(auc_url)}</a>"
                            )
                        })
                        state["notifiedAuction1h"] = auc_1h_key
                        state_changed = True

            # 4. Завершен прием заявок -> переход к ожиданию аукциона (active.auction)
            if t_status == "active.auction" and state.get("notifiedStage") != "stage_active_auction":
                auc_time_str = format_deadline_strict(auc_start_iso) if auc_start_iso else "определяется системой Prozorro"
                part_lines = []
                for idx, p in enumerate(participants, start=1):
                    part_lines.append(
                        f"{idx}) <b>{html.escape(p['name'])}</b> — {html.escape(p['priceFormatted'])}"
                    )
                if not part_lines:
                    part_lines.append("<i>Список участников формируется системой Prozorro (названия скрыты до завершения аукциона).</i>")

                events_for_all.append({
                    "title": "🔨 ЗАВЕРШЕН ПРИЕМ ЗАЯВОК: ОЖИДАНИЕ ЭЛЕКТРОННОГО АУКЦИОНА",
                    "bodyHtml": (
                        f"• <b>Дальнейший статус закупки:</b> Ожидание электронного аукциона\n"
                        f"• <b>Дата и время старта аукциона:</b> <b>{html.escape(auc_time_str)}</b>\n\n"
                        f"<b>Заявившиеся участники, ожидающие аукциона:</b>\n" + "\n".join(part_lines)
                    )
                })
                state["notifiedStage"] = "stage_active_auction"
                state["liveMenuStatus"] = f"🔨 Ожидание электронного аукциона (старт: {auc_time_str})"
                state_changed = True

            # 5. Контроль дисквалификации лидера / любого участника
            rejected_seen = set(state.get("rejectedAwardsSeen") or [])
            pending_awards = [a for a in awards if a.get("status") == "pending"]
            active_awards = [a for a in awards if a.get("status") == "active"]

            for aw in awards:
                aw_id = aw.get("id", "")
                if aw.get("status") == "unsuccessful" and aw_id and aw_id not in rejected_seen:
                    supp = (aw.get("suppliers") or [{}])[0]
                    rej_name = supp.get("name", "Участник")
                    rej_edrpou = str(supp.get("identifier", {}).get("id", "Не указан"))
                    reason_summary = summarize_rejection_reason_with_gemini(tender_ua_id, rej_name, aw)

                    if pending_awards:
                        next_aw = pending_awards[-1]
                        nsupp = (next_aw.get("suppliers") or [{}])[0]
                        next_name = nsupp.get("name", "Следующий участник")
                        next_edrpou = str(nsupp.get("identifier", {}).get("id", "Не указан"))
                        nval = next_aw.get("value", {})
                        next_price = format_participant_price_with_vat(
                            float(nval.get("amount", 0) or 0),
                            bool(nval.get("valueAddedTaxIncluded", False)),
                            is_zero_vat, items_count, single_qty, unit_label
                        )
                        transition_line = (
                            f"<b>Право заключения договора (оценка предложения) перешло к следующему участнику:</b>\n"
                            f"<b>{html.escape(next_name)}</b> (ЕГРПОУ: {html.escape(next_edrpou)}) — <b>{html.escape(next_price)}</b>"
                        )
                    elif t_status in ("unsuccessful", "cancelled"):
                        transition_line = "<b>Дальнейший статус:</b> Все предложения отклонены — торги не состоялись / отменены."
                    else:
                        transition_line = "<b>Дальнейший статус:</b> Ожидается публикация решения по следующему участнику."

                    events_for_all.append({
                        "title": "⚠️ ВНИМАНИЕ! УЧАСТНИК ТОРГОВ ОТКЛОНЕН ЗАКАЗЧИКОМ",
                        "bodyHtml": (
                            f"• <b>Отклоненный участник:</b> <b>{html.escape(rej_name)}</b> (ЕГРПОУ: {html.escape(rej_edrpou)})\n"
                            f"• <b>Причина отклонения:</b> <i>{html.escape(reason_summary)}</i>\n\n"
                            f"{transition_line}"
                        )
                    })
                    rejected_seen.add(aw_id)
                    state["rejectedAwardsSeen"] = list(rejected_seen)
                    state_changed = True

            # 6. Оценка предложений заказчиком (active.qualification / pending award)
            if pending_awards:
                cur_aw = pending_awards[-1]
                cur_aw_id = cur_aw.get("id", "")
                supp = (cur_aw.get("suppliers") or [{}])[0]
                leader_name = supp.get("name", "Не указано")
                leader_edrpou = str(supp.get("identifier", {}).get("id", "Не указан"))
                val = cur_aw.get("value", {})
                leader_price = format_participant_price_with_vat(
                    float(val.get("amount", 0) or 0),
                    bool(val.get("valueAddedTaxIncluded", False)),
                    is_zero_vat, items_count, single_qty, unit_label
                )

                stage_key = f"pending_award_{cur_aw_id}"
                if state.get("notifiedStage") != stage_key:
                    events_for_all.append({
                        "title": "📊 ОЦЕНКА ПРЕДЛОЖЕНИЙ (КВАЛИФИКАЦИЯ)",
                        "bodyHtml": (
                            f"• <b>Дальнейший статус закупки:</b> Оценка предложений заказчиком\n"
                            f"• <b>На рассмотрении сейчас находится предложение:</b> <b>{html.escape(leader_name)}</b> (ЕГРПОУ: {html.escape(leader_edrpou)}) — <b>{html.escape(leader_price)}</b>"
                        ),
                        "includeParticipants": True
                    })
                    state["notifiedStage"] = stage_key
                    state["liveMenuStatus"] = f"⏳ Оценка предложения заказчиком — рассматривается {leader_name} ({leader_price})"
                    state_changed = True

            # 7. Уведомление Админу о предоставлении 24 часов на исправление ошибок
            seen_24h = set(state.get("seen24hIds") or [])
            for aw in awards:
                supp = (aw.get("suppliers") or [{}])[0]
                p_name = supp.get("name", "Участник")
                p_edrpou = str(supp.get("identifier", {}).get("id", "Не указан"))
                for ms in (aw.get("milestones") or []):
                    if ms.get("code") == "24h":
                        ms_id = ms.get("id") or f"{aw.get('id')}_{ms.get('dueDate')}"
                        due_str = format_deadline_strict(ms.get("dueDate", ""))
                        desc_24h = ms.get("description", "Устранение несоответствий в документах тендерного предложения")
                        if ms_id not in seen_24h:
                            events_for_admin.append({
                                "title": "⏱ ЗАКАЗЧИКОМ ДАНО 24 ЧАСА НА ИСПРАВЛЕНИЕ ОШИБОК",
                                "bodyHtml": (
                                    f"• <b>Участник:</b> <b>{html.escape(p_name)}</b> (ЕГРПОУ: {html.escape(p_edrpou)})\n"
                                    f"• <b>Исправить недостатки в срок до:</b> <b>{html.escape(due_str)}</b>\n"
                                    f"• <b>Причина / перечень замечаний:</b> <i>{html.escape(desc_24h)}</i>"
                                )
                            })
                            seen_24h.add(ms_id)
                            state["seen24hIds"] = list(seen_24h)
                            state["liveMenuStatus"] = (
                                f"⏱ Оценка предложения {p_name} — заказчиком дано 24 часа на исправление недостатков "
                                f"тендерного предложения в срок до {due_str}"
                            )
                            state_changed = True

            # 8. Принятие заказчиком решения о намерении заключить договор (Акцепт)
            if active_awards:
                win_aw = active_awards[-1]
                win_aw_id = win_aw.get("id", "")
                supp = (win_aw.get("suppliers") or [{}])[0]
                win_name = supp.get("name", "Не указано")
                win_edrpou = str(supp.get("identifier", {}).get("id", "Не указан"))
                val = win_aw.get("value", {})
                win_price = format_participant_price_with_vat(
                    float(val.get("amount", 0) or 0),
                    bool(val.get("valueAddedTaxIncluded", False)),
                    is_zero_vat, items_count, single_qty, unit_label
                )
                stage_key = f"active_award_{win_aw_id}"
                if state.get("notifiedStage") != stage_key:
                    events_for_all.append({
                        "title": "🏆 ПРИНЯТО РЕШЕНИЕ О НАМЕРЕНИИ ЗАКЛЮЧИТЬ ДОГОВОР (АКЦЕПТ)",
                        "bodyHtml": (
                            f"• <b>Победитель:</b> <b>{html.escape(win_name)}</b> (ЕГРПОУ: {html.escape(win_edrpou)})\n"
                            f"• <b>Ценовое предложение победителя:</b> <b>{html.escape(win_price)}</b>\n"
                            f"• <b>Статус участника:</b> Признан победителем, ожидается заключение договора"
                        ),
                        "includeParticipants": True,
                        "winnerSummary": f"{win_name} ({win_price})"
                    })
                    state["notifiedStage"] = stage_key
                    state["winnerName"] = win_name
                    state["liveMenuStatus"] = f"🏆 {win_name} признан победителем — ожидается заключение договора"
                    state_changed = True

            # 9. Подписание договора, отмена тендера или несостоявшиеся торги -> АВТООЧИСТКА из меню «Принимаем участие»
            active_contracts = [c for c in contracts if c.get("status") == "active"]
            if active_contracts or t_status == "complete":
                c_obj = active_contracts[-1] if active_contracts else (contracts[-1] if contracts else {})
                date_signed = format_date_only(c_obj.get("dateSigned") or c_obj.get("dateModified") or "")
                win_name = state.get("winnerName") or ""
                if not win_name and active_awards:
                    win_name = ((active_awards[-1].get("suppliers") or [{}])[0]).get("name", "Победитель")

                if state.get("notifiedStage") != "stage_contract_signed":
                    events_for_all.append({
                        "title": "🤝 ЗАКЛЮЧЕН ДОГОВОР ПО ТЕНДЕРУ",
                        "bodyHtml": (
                            f"• <b>Победитель:</b> <b>{html.escape(win_name or 'Определен')}</b>\n"
                            f"• <b>Дата заключения договора:</b> <b>{html.escape(date_signed or 'Зарегистрирован в Prozorro')}</b>\n"
                            f"• <b>Статус в меню:</b> Тендер завершен и автоматически убран из активного списка «Принимаем участие»."
                        ),
                        "closeParticipationMenu": "Договор заключен",
                        "winnerName": win_name,
                        "contractDate": date_signed
                    })
                    state["notifiedStage"] = "stage_contract_signed"
                    state["winnerName"] = win_name
                    state["contractDate"] = date_signed
                    state_changed = True

            elif t_status == "unsuccessful" and state.get("notifiedStage") != "stage_unsuccessful":
                reason_unsucc = (
                    "На торги не заявился ни один участник."
                    if not participants
                    else "Отклонены все поданные тендерные предложения."
                )
                events_for_all.append({
                    "title": "❌ ТЕНДЕР НЕ СОСТОЯЛСЯ",
                    "bodyHtml": (
                        f"• <b>Дальнейший статус закупки:</b> Торги не состоялись\n"
                        f"• <b>Причина:</b> {html.escape(reason_unsucc)}\n"
                        f"• <b>Статус в меню:</b> Тендер автоматически убран из активного списка «Принимаем участие»."
                    ),
                    "closeParticipationMenu": "Торги не состоялись"
                })
                state["notifiedStage"] = "stage_unsuccessful"
                state_changed = True

            elif t_status == "cancelled" and state.get("notifiedStage") != "stage_cancelled":
                canc_list = t_data.get("cancellations", []) or []
                canc_reason = canc_list[-1].get("reason", "Решение заказчика в Prozorro") if canc_list else "Решение заказчика"
                events_for_all.append({
                    "title": "🚫 ТЕНДЕР ОТМЕНЕН ЗАКАЗЧИКОМ",
                    "bodyHtml": (
                        f"• <b>Причина отмены:</b> <i>{html.escape(canc_reason)}</i>\n"
                        f"• <b>Статус в меню:</b> Тендер автоматически убран из активного списка «Принимаем участие»."
                    ),
                    "closeParticipationMenu": "Торги отменены"
                })
                state["notifiedStage"] = "stage_cancelled"
                state_changed = True

            # 10. Вопросы заказчику, требования/жалобы в АМКУ и мониторинг ДАСУ (только для Админа)
            seen_questions = set(state.get("seenQuestions") or [])
            for q in (t_data.get("questions") or []):
                q_id = q.get("id", "")
                q_title = q.get("title", "")
                q_desc = q.get("description", "")
                q_ans = q.get("answer", "")
                q_key = f"{q_id}_{'ans' if q_ans else 'ask'}"
                if q_id and q_key not in seen_questions:
                    body_q = f"• <b>Обращение/вопрос:</b> <i>{html.escape(q_title)} — {html.escape(q_desc[:350])}</i>"
                    if q_ans:
                        body_q += f"\n• <b>Ответ заказчика:</b> <i>{html.escape(q_ans[:400])}</i>"
                    events_for_admin.append({
                        "title": "💬 ВОПРОС К ЗАКАЗЧИКУ В PROZORRO" if not q_ans else "💬 ОПУБЛИКОВАН ОТВЕТ ЗАКАЗЧИКА НА ВОПРОС",
                        "bodyHtml": body_q
                    })
                    seen_questions.add(q_key)
                    state["seenQuestions"] = list(seen_questions)
                    state_changed = True

            seen_complaints = set(state.get("seenComplaints") or [])
            all_complaints = list(t_data.get("complaints") or [])
            for aw in awards:
                all_complaints.extend(aw.get("complaints") or [])
            for comp in all_complaints:
                c_id = comp.get("id", "")
                c_status = comp.get("status", "")
                if c_status == "draft":
                    continue
                c_key = f"{c_id}_{c_status}_{comp.get('ResolutionType', '')}"
                if c_id and c_key not in seen_complaints:
                    c_type_label = "ЖАЛОБА В АМКУ" if comp.get("type") == "complaint" else "ТРЕБОВАНИЕ К ЗАКАЗЧИКУ"
                    c_author = (comp.get("author") or {}).get("name", "Участник")
                    c_title = comp.get("title", "")
                    c_desc = comp.get("description", "")
                    c_res = comp.get("resolution", "") or comp.get("tendererAction", "")
                    body_c = (
                        f"• <b>Податель:</b> {html.escape(c_author)}\n"
                        f"• <b>Статус:</b> <b>{html.escape(c_status)}</b>\n"
                        f"• <b>Суть / Причина:</b> <i>{html.escape(c_title)} — {html.escape(c_desc[:400])}</i>"
                    )
                    if c_res:
                        body_c += f"\n• <b>Решение / Ответ:</b> <i>{html.escape(c_res[:400])}</i>"
                    events_for_admin.append({
                        "title": f"⚖️ МОНИТОРИНГ АМКУ / ТРЕБОВАНИЙ: {c_type_label}",
                        "bodyHtml": body_c
                    })
                    seen_complaints.add(c_key)
                    state["seenComplaints"] = list(seen_complaints)
                    state_changed = True

            if lookup_id in dasu_map:
                m_data = dasu_map[lookup_id]
                m_status = m_data.get("status", "")
                posts = m_data.get("posts", []) or []
                conclusion = m_data.get("conclusion", {}) or {}
                dasu_key = f"{m_data.get('id')}_{m_status}_{len(posts)}_{bool(conclusion)}"
                if state.get("seenDasuKey") != dasu_key:
                    reasons = ", ".join(m_data.get("reasons", []) or ["Основания ст. 8 Закона о публичных закупках"])
                    body_d = f"• <b>Статус мониторинга ДАСУ:</b> <b>{html.escape(m_status)}</b>\n• <b>Основание:</b> {html.escape(reasons)}"
                    if posts:
                        last_p = posts[-1]
                        p_party = "Ответ заказчика на запрос ДАСУ" if last_p.get("author") == "tender_owner" else "Запрос ДАСУ заказчику"
                        body_d += f"\n• <b>{p_party}:</b> <i>{html.escape(last_p.get('title', ''))} — {html.escape(last_p.get('description', '')[:350])}</i>"
                    if conclusion:
                        viol = "Выявлены нарушения" if conclusion.get("violationOccurred") else "Нарушений не выявлено"
                        body_d += f"\n• <b>Заключение ДАСУ:</b> <b>{viol}</b> (<i>{html.escape(str(conclusion.get('description', ''))[:350])}</i>)"
                    events_for_admin.append({
                        "title": "🏛 МОНИТОРИНГ ГОСАУДИТСЛУЖБЫ (ДАСУ)",
                        "bodyHtml": body_d
                    })
                    state["seenDasuKey"] = dasu_key
                    state_changed = True

            if events_for_all or events_for_admin or state_changed:
                update_payload = {
                    "source": "python_agent",
                    "secret": SECRET_BRIDGE_KEY,
                    "action": "process_tracking_events",
                    "tenderID": tender_ua_id,
                    "itemsCount": items_count,
                    "participants": participants,
                    "eventsForAll": events_for_all,
                    "eventsForAdmin": events_for_admin,
                    "trackingState": state
                }
                http_session.post(GOOGLE_WEB_APP_URL, json=update_payload, timeout=25)

        # Фоновое обогащение таблицы (Победитель + Дата договора по всем тендерам таблицы для кнопки «✒️ База данных», $0.00 без ИИ)
        bg_updates = []
        for bg in bg_list[:15]:
            bg_ua_id = bg.get("tenderID", "")
            bg_int_id = bg.get("prozorroInternalId") or bg_ua_id
            t_bg = fetch_tender_detail(bg_int_id)
            if not t_bg:
                continue
            bg_status = t_bg.get("status", "")
            bg_contracts = [c for c in (t_bg.get("contracts") or []) if c.get("status") == "active"]
            bg_awards = [a for a in (t_bg.get("awards") or []) if a.get("status") == "active"]

            if bg_contracts or bg_status == "complete":
                c_obj = bg_contracts[-1] if bg_contracts else ((t_bg.get("contracts") or [{}])[-1])
                c_date = format_date_only(c_obj.get("dateSigned") or c_obj.get("dateModified") or "")
                w_name = ""
                if bg_awards:
                    w_name = ((bg_awards[-1].get("suppliers") or [{}])[0]).get("name", "")
                if w_name:
                    bg_updates.append({
                        "yearStr": bg.get("yearStr"),
                        "tenderID": bg_ua_id,
                        "winnerName": w_name,
                        "contractDate": c_date or "Подписан"
                    })

        if bg_updates:
            http_session.post(GOOGLE_WEB_APP_URL, json={
                "source": "python_agent",
                "secret": SECRET_BRIDGE_KEY,
                "action": "save_background_history",
                "updates": bg_updates
            }, timeout=20)

    except Exception as e:
        print(f"Ошибка при проверке сопровождения тендеров: {e}")

def sync_stats_to_crm(gemini_used_this_run: int):
    try:
        start_date = get_sync_val("monitoring_start_date")
        if not start_date:
            start_date = "20.10.2026"
            set_sync_val("monitoring_start_date", start_date)

        total_taken_db = get_total_processed_count()
        prev_ai_str = get_sync_val("stats_ai_analyzed")
        prev_ai = int(prev_ai_str) if prev_ai_str.isdigit() else 28
        total_ai_analyzed = prev_ai + gemini_used_this_run
        set_sync_val("stats_ai_analyzed", str(total_ai_analyzed))

        http_session.post(GOOGLE_WEB_APP_URL, json={
            "source": "python_agent",
            "secret": SECRET_BRIDGE_KEY,
            "action": "sync_stats",
            "monitoringStartDate": start_date,
            "totalTaken": total_taken_db,
            "totalAiAnalyzed": total_ai_analyzed
        }, timeout=15)
    except Exception as e:
        print(f"Ошибка синхронизации статистики: {e}")

def fetch_tender_detail(t_id: str) -> dict:
    for attempt in range(4):
        try:
            r = http_session.get(f"{PROZORRO_API_BASE}/tenders/{t_id}", timeout=12)
            if r.status_code == 200:
                return r.json().get("data", {})
            elif r.status_code == 429:
                time.sleep(1.5 * (attempt + 1))
                continue
        except Exception:
            pass
    return {}

def evaluate_active_tenders(active_ids: list, gemini_budget_left: int) -> tuple[int, int, int, int]:
    matched_budget = 0
    matched_fuel_cpv = 0
    sent_to_crm = 0
    gemini_used = 0
    fuel_tenders_to_analyze = []

    with ThreadPoolExecutor(max_workers=15) as executor:
        future_to_id = {executor.submit(fetch_tender_detail, tid): tid for tid in active_ids}
        for future in as_completed(future_to_id):
            tid = future_to_id[future]
            tender = future.result()
            if not tender:
                add_to_retry_queue(tid)
                continue

            status = tender.get("status", "")
            if status not in ("active.enquiries", "active.tendering"):
                mark_as_processed(tid)
                continue

            amount = float(tender.get("value", {}).get("amount", 0) or 0)
            if amount < MIN_AMOUNT_UAH:
                mark_as_processed(tid)
                continue

            matched_budget += 1

            is_valid_fuel, _ = is_valid_diesel_or_gasoline_tender(tender)
            if not is_valid_fuel:
                mark_as_processed(tid)
                continue

            matched_fuel_cpv += 1
            fuel_tenders_to_analyze.append(tender)

    for tender in fuel_tenders_to_analyze:
        t_id = tender["id"]
        tender_ua_id = tender.get("tenderID", t_id)
        amount = float(tender.get("value", {}).get("amount", 0) or 0)
        print(f"Целевой тендер (Бензин/ДТ >= 2.4 млн): {tender_ua_id} | {amount:,.2f} грн | {tender.get('title', '')[:70]}")

        if is_strictly_coupons_or_cards(tender):
            print(f" -> Мгновенно отсеян (явные талоны/топливные карты без налива): {tender_ua_id}")
            mark_as_processed(t_id)
            mark_as_processed(tender_ua_id)
            continue

        if gemini_used >= gemini_budget_left:
            add_to_retry_queue(t_id)
            continue

        try:
            if gemini_used > 0:
                time.sleep(2)
            print(f" -> Скачивание документации и экспертиза Gemini для {tender_ua_id}...")
            docs_text, pdf_parts = extract_all_attachments(tender.get("documents", []))
            ai_verdict = analyze_tender_with_gemini(tender, docs_text, pdf_parts)
            gemini_used += 1

            if ai_verdict.get("is_relevant"):
                print(f" -> ПОДТВЕРЖДЕН НАЛИВ! Отправка в Google Sheets и Telegram: {tender_ua_id}")
                dispatch_to_crm(tender, ai_verdict, docs_text)
                sent_to_crm += 1
            else:
                print(f" -> Отклонен ИИ после чтения ТД: {tender_ua_id}")

            mark_as_processed(t_id)
            mark_as_processed(tender_ua_id)
        except Exception as e:
            gemini_used += 1
            print(f" -> Ошибка при анализе {tender_ua_id} (сохранен в retry_queue): {e}")
            add_to_retry_queue(t_id)

    return matched_budget, matched_fuel_cpv, sent_to_crm, gemini_used

def run_monitoring():
    total_scanned = 0
    total_active = 0
    total_budget = 0
    total_fuel = 0
    total_sent = 0
    total_gemini_used = 0
    gemini_left = MAX_GEMINI_ANALYSES_PER_RUN

    retry_ids = get_retry_queue_ids()
    if retry_ids:
        rb, rf, rs, g_used = evaluate_active_tenders(retry_ids, gemini_left)
        gemini_left -= g_used
        total_gemini_used += g_used
        total_active += len(retry_ids)
        total_budget += rb
        total_fuel += rf
        total_sent += rs

    last_top_ts = get_sync_val("last_top_ts")
    new_top_ts = ""
    next_url = f"{PROZORRO_API_BASE}/tenders?descending=1&limit=100&opt_fields=status"
    fresh_active_ids = []

    print("=== Проверка свежих изменений ленты Prozorro ===")
    for page in range(1, 16):
        if not next_url:
            break
        resp = http_session.get(next_url, timeout=15).json()
        items_list = resp.get("data", [])
        if not items_list:
            break

        if page == 1 and items_list:
            new_top_ts = items_list[0].get("dateModified", "")

        reached_previous_window = False
        for item in items_list:
            total_scanned += 1
            t_id = item.get("id")
            status = item.get("status", "")
            mod_ts = item.get("dateModified", "")

            if status in ("active.enquiries", "active.tendering") and not is_already_processed(t_id):
                fresh_active_ids.append(t_id)

            if last_top_ts and mod_ts < last_top_ts:
                reached_previous_window = True

        next_url = resp.get("next_page", {}).get("uri")
        if reached_previous_window:
            print(f"Окно синхронизации сомкнуто со штампом прошлого запуска ({last_top_ts}) на стр. {page}.")
            break

    if new_top_ts:
        set_sync_val("last_top_ts", new_top_ts)

    print(f"Найдено свежих активных тендеров: {len(fresh_active_ids)}")
    b1, f1, s1, g1 = evaluate_active_tenders(fresh_active_ids, gemini_left)
    total_gemini_used += g1
    total_active += len(fresh_active_ids)
    total_budget += b1
    total_fuel += f1
    total_sent += s1

    check_tracked_and_background_tenders()
    sync_stats_to_crm(total_gemini_used)

    remaining_in_queue = len(get_retry_queue_ids())
    print("=========================================")
    print("ИТОГИ ЦИКЛА МОНИТОРИНГА:")
    print(f"- Всего просмотрено записей в ленте: {total_scanned}")
    print(f"- Из них проверено активных карточек: {total_active}")
    print(f"- Активных тендеров с бюджетом >= 2.4 млн грн: {total_budget}")
    print(f"- Из них целевых (бензин / ДТ): {total_fuel}")
    print(f"- Подтверждено наливом и отправлено в Telegram: {total_sent}")
    print(f"- Осталось в очереди retry_queue на следующий цикл: {remaining_in_queue}")
    print("=========================================")

if __name__ == "__main__":
    run_monitoring()
